from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .audit_models import AuditSeverity, MetadataAuditIssue
from .constants import DEFAULT_REPORT_COLORS, EXPORT_COLUMNS
from .models import Track, TrackStatus

ORANGE = "E97824"
DARK_ORANGE = "C85A0A"
WARM_WHITE = "FFFDF9"
LIGHT_ORANGE = "FFF0E3"
GRAPHITE = "4B5158"
METAL = "D1D5DB"
TEXT = "2F3439"
MUTED = "737980"
BLUE = "E2EFFB"
BLUE_DARK = "2F75B5"
GREEN = "E5F3E8"
GREEN_DARK = "2E7D32"
CORAL = "F8DFDA"
CORAL_DARK = "A9473E"
ROW_BORDER = Side(style="thin", color="E2DED8")
HEADER_BORDER = Side(style="medium", color=ORANGE)
DATA_ROW_HEIGHT = 17.25
HEADER_ROW = 4
DATA_START_ROW = 5

# openpyxl style objects are immutable and interned per workbook, so sharing
# one instance across thousands of cells is both safe and much cheaper than
# allocating a new Font/Border/Alignment for every cell.
DATA_FONT = Font(color=TEXT, size=9)
DATA_BORDER = Border(bottom=ROW_BORDER)
DATA_ALIGNMENT = Alignment(vertical="center")
CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
_FILL_CACHE: dict[str, PatternFill] = {}


def _fill(color: str) -> PatternFill:
    fill = _FILL_CACHE.get(color)
    if fill is None:
        fill = _FILL_CACHE[color] = PatternFill("solid", fgColor=color)
    return fill


def _safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    # One control character in any tag must not abort the whole report:
    # openpyxl raises IllegalCharacterError for them.
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _color(value: str | None, fallback: str) -> str:
    normalized = (value or "").lstrip("#").upper()
    if len(normalized) == 6 and all(
        character in "0123456789ABCDEF" for character in normalized
    ):
        return normalized
    return fallback


def _status_colors(colors: dict[str, str]) -> dict[str, str]:
    return {
        status: _color(
            colors.get(status),
            DEFAULT_REPORT_COLORS.get(status, WARM_WHITE),
        )
        for status in DEFAULT_REPORT_COLORS
    }


def _summary_text(tracks: list[Track]) -> str:
    counts = Counter(track.status for track in tracks)
    issue_count = len(
        {track.collection_issue_id for track in tracks if track.collection_issue_id}
    )
    return (
        f"Файлов: {len(tracks)}  ·  "
        f"Уникальных: {counts[TrackStatus.UNIQUE]}  ·  "
        f"Дубликатов: {counts[TrackStatus.DUPLICATE]}  ·  "
        f"Вариаций: {counts[TrackStatus.VARIATION]}  ·  "
        f"Конфликтов: {counts[TrackStatus.CONFLICT]}  ·  "
        f"Предупреждений по комплектам: {issue_count}"
    )


def _prepare_sheet(
    sheet,
    title: str,
    subtitle: str,
    headers: list[str],
    *,
    freeze_panes: str = "A5",
    tab_color: str = ORANGE,
) -> None:
    last_column = len(headers)
    last_letter = get_column_letter(last_column)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = tab_color
    sheet.freeze_panes = freeze_panes
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_letter}{HEADER_ROW}"

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = sheet.cell(1, 1, title)
    title_cell.fill = PatternFill("solid", fgColor=ORANGE)
    title_cell.font = Font(color="FFFFFF", bold=True, size=15)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28
    for column in range(1, last_column + 1):
        sheet.cell(1, column).fill = PatternFill("solid", fgColor=ORANGE)

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    subtitle_cell = sheet.cell(2, 1, subtitle)
    subtitle_cell.fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
    subtitle_cell.font = Font(color=MUTED, italic=True, size=9)
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 19
    for column in range(1, last_column + 1):
        sheet.cell(2, column).fill = PatternFill("solid", fgColor=LIGHT_ORANGE)

    sheet.row_dimensions[3].height = 7
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(HEADER_ROW, column, header)
        cell.fill = PatternFill("solid", fgColor=GRAPHITE)
        cell.font = Font(color="FFFFFF", bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=HEADER_BORDER)
    sheet.row_dimensions[HEADER_ROW].height = 22


def _style_data_row(
    sheet,
    row_number: int,
    column_count: int,
    fill_color: str = WARM_WHITE,
) -> None:
    # Address cells explicitly: ``sheet[row_number]`` asks openpyxl for
    # ``max_column``, which scans every cell of the sheet and turns row
    # styling into an O(rows²) operation on large reports.
    sheet.row_dimensions[row_number].height = DATA_ROW_HEIGHT
    fill = _fill(fill_color)
    for column in range(1, column_count + 1):
        cell = sheet.cell(row_number, column)
        cell.fill = fill
        cell.font = DATA_FONT
        cell.border = DATA_BORDER
        cell.alignment = DATA_ALIGNMENT


def _fit_columns(
    sheet,
    *,
    header_row: int = HEADER_ROW,
    maximum: int = 48,
    minimum: int = 10,
) -> None:
    # ``max_row``/``max_column`` are full scans in openpyxl; evaluate them once.
    last_row = sheet.max_row
    last_column = sheet.max_column
    for column in range(1, last_column + 1):
        length = len(str(sheet.cell(header_row, column).value or ""))
        for row in range(header_row + 1, last_row + 1):
            value = sheet.cell(row, column).value
            if value is not None:
                line_lengths = [len(part) for part in str(value).splitlines()]
                length = max(length, max(line_lengths, default=0))
        sheet.column_dimensions[get_column_letter(column)].width = min(
            max(length + 2, minimum),
            maximum,
        )


def _finish_table(sheet) -> None:
    last_letter = get_column_letter(sheet.max_column)
    sheet.auto_filter.ref = (
        f"A{HEADER_ROW}:{last_letter}{max(sheet.max_row, HEADER_ROW)}"
    )


def _write_track_sheet(
    workbook: Workbook,
    name: str,
    tracks: list[Track],
    fields: list[str],
    colors: dict[str, str],
) -> None:
    sheet = workbook.create_sheet(name)
    headers = ["№"] + [EXPORT_COLUMNS[field][0] for field in fields]
    _prepare_sheet(sheet, name, _summary_text(tracks), headers)
    status_colors = _status_colors(colors)
    warning_column = (
        fields.index("collection_warning") + 2
        if "collection_warning" in fields
        else None
    )
    warning_font = Font(color=CORAL_DARK, bold=True, size=9)
    for number, track in enumerate(tracks, start=1):
        row_number = DATA_START_ROW + number - 1
        values = [number] + [_safe(value) for value in track.table_row(fields)]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        _style_data_row(
            sheet,
            row_number,
            len(headers),
            status_colors.get(track.status.value, WARM_WHITE),
        )
        sheet.cell(row_number, 1).alignment = CENTER_ALIGNMENT
        if warning_column and track.collection_warning:
            warning_cell = sheet.cell(row_number, warning_column)
            warning_cell.fill = _fill(CORAL)
            warning_cell.font = warning_font

    sheet.column_dimensions["A"].width = 7
    _fit_columns(sheet)
    _finish_table(sheet)


def _write_matrix(
    workbook: Workbook,
    tracks: list[Track],
    colors: dict[str, str],
) -> bool:
    attention = [
        track
        for track in tracks
        if track.status
        in (TrackStatus.DUPLICATE, TrackStatus.VARIATION, TrackStatus.CONFLICT)
    ]
    if not attention:
        return False
    folders = sorted({track.relative_folder for track in attention}, key=str.casefold)
    titles = sorted({track.base_title for track in attention}, key=str.casefold)
    sheet = workbook.create_sheet("Матрица дубликатов")
    headers = ["Базовое имя", *folders]
    _prepare_sheet(
        sheet,
        "Матрица дубликатов и вариаций",
        (
            f"Связанных файлов: {len(attention)}  ·  "
            f"Названий: {len(titles)}  ·  Папок: {len(folders)}"
        ),
        headers,
        freeze_panes="B5",
        tab_color=BLUE_DARK,
    )
    status_colors = _status_colors(colors)
    cells: dict[tuple[str, str], list[tuple[TrackStatus, str]]] = defaultdict(list)
    for track in attention:
        if track.status == TrackStatus.VARIATION:
            marker = f"Вариация: {track.variation or 'версия'} — {track.filename}"
        elif track.status == TrackStatus.CONFLICT:
            marker = f"Конфликт — {track.filename}"
        else:
            marker = f"Дубликат — {track.filename}"
        cells[(track.base_title, track.relative_folder)].append(
            (track.status, marker)
        )
    title_font = Font(color=TEXT, bold=True, size=9)
    for offset, title in enumerate(titles):
        row_number = DATA_START_ROW + offset
        sheet.cell(row_number, 1, _safe(title))
        _style_data_row(sheet, row_number, len(headers))
        sheet.cell(row_number, 1).fill = _fill(LIGHT_ORANGE)
        sheet.cell(row_number, 1).font = title_font
        for column, folder in enumerate(folders, start=2):
            items = cells.get((title, folder))
            if not items:
                continue
            sheet.cell(
                row_number,
                column,
                _safe("; ".join(marker for _status, marker in items)),
            )
            statuses = {status for status, _marker in items}
            if TrackStatus.CONFLICT in statuses:
                fill_color = status_colors[TrackStatus.CONFLICT.value]
            elif TrackStatus.VARIATION in statuses:
                fill_color = status_colors[TrackStatus.VARIATION.value]
            else:
                fill_color = status_colors[TrackStatus.DUPLICATE.value]
            sheet.cell(row_number, column).fill = _fill(fill_color)

    _fit_columns(sheet, maximum=55, minimum=14)
    _finish_table(sheet)
    return True


def _write_folder_summary(
    workbook: Workbook,
    tracks: list[Track],
    colors: dict[str, str],
) -> bool:
    if not tracks:
        return False
    grouped: dict[str, list[Track]] = defaultdict(list)
    for track in tracks:
        grouped[track.relative_folder].append(track)
    sheet = workbook.create_sheet("Сводка по папкам")
    headers = [
        "Папка",
        "Всего файлов",
        "Уникальных",
        "Дубликатов",
        "Вариаций",
        "Конфликтов",
        "Форматы",
        "Размер (MB)",
        "Проверка комплекта",
    ]
    issue_count = len(
        {track.collection_issue_id for track in tracks if track.collection_issue_id}
    )
    _prepare_sheet(
        sheet,
        "Сводка по папкам",
        (
            f"Папок: {len(grouped)}  ·  Файлов: {len(tracks)}  ·  "
            f"Папок/альбомов с предупреждениями: {issue_count}"
        ),
        headers,
        tab_color=GREEN_DARK,
    )
    status_colors = _status_colors(colors)
    warning_font = Font(color=CORAL_DARK, bold=True, size=9)
    ok_font = Font(color=GREEN_DARK, bold=True, size=9)
    for offset, folder in enumerate(sorted(grouped, key=str.casefold)):
        rows = grouped[folder]
        counts = Counter(track.status for track in rows)
        warnings = " · ".join(
            sorted(
                {
                    track.collection_warning
                    for track in rows
                    if track.collection_warning
                }
            )
        )
        values = [
            _safe(folder),
            len(rows),
            counts[TrackStatus.UNIQUE],
            counts[TrackStatus.DUPLICATE],
            counts[TrackStatus.VARIATION],
            counts[TrackStatus.CONFLICT],
            ", ".join(sorted({track.format for track in rows})),
            round(sum(track.size_mb for track in rows), 2),
            warnings or "Комплект без обнаруженных пропусков",
        ]
        row_number = DATA_START_ROW + offset
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        _style_data_row(sheet, row_number, len(headers))
        sheet.cell(row_number, 3).fill = _fill(status_colors[TrackStatus.UNIQUE.value])
        sheet.cell(row_number, 4).fill = _fill(
            status_colors[TrackStatus.DUPLICATE.value]
        )
        sheet.cell(row_number, 5).fill = _fill(
            status_colors[TrackStatus.VARIATION.value]
        )
        sheet.cell(row_number, 6).fill = _fill(status_colors[TrackStatus.CONFLICT.value])
        completeness_cell = sheet.cell(row_number, 9)
        completeness_cell.fill = _fill(CORAL if warnings else GREEN)
        completeness_cell.font = warning_font if warnings else ok_font

    counts = Counter(track.status for track in tracks)
    total_row = DATA_START_ROW + len(grouped)
    total_values = [
        "ИТОГО",
        len(tracks),
        counts[TrackStatus.UNIQUE],
        counts[TrackStatus.DUPLICATE],
        counts[TrackStatus.VARIATION],
        counts[TrackStatus.CONFLICT],
        ", ".join(sorted({track.format for track in tracks})),
        round(sum(track.size_mb for track in tracks), 2),
        f"Папок/альбомов с предупреждениями: {issue_count}",
    ]
    for column, value in enumerate(total_values, start=1):
        sheet.cell(total_row, column, value)
    _style_data_row(sheet, total_row, len(headers), LIGHT_ORANGE)
    total_font = Font(color=DARK_ORANGE, bold=True, size=9)
    total_border = Border(top=HEADER_BORDER, bottom=ROW_BORDER)
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(total_row, column)
        cell.font = total_font
        cell.border = total_border

    _fit_columns(sheet, maximum=52)
    _finish_table(sheet)
    return True


def _write_audit_sheet(
    workbook: Workbook,
    issues: list[MetadataAuditIssue],
    tracks_by_id: dict[str, Track],
) -> None:
    sheet = workbook.create_sheet("Аудит метаданных")
    headers = [
        "№",
        "Уровень",
        "Категория",
        "Поле",
        "Описание",
        "Текущее значение",
        "Предложение",
        "Уверенность",
        "Автоисправление",
        "Папка",
        "Файл",
        "Полный путь",
    ]
    counts = Counter(issue.severity for issue in issues)
    subtitle = (
        f"Замечаний: {len(issues)}  ·  "
        f"Ошибок: {counts[AuditSeverity.ERROR]}  ·  "
        f"Предупреждений: {counts[AuditSeverity.WARNING]}  ·  "
        f"Рекомендаций: {counts[AuditSeverity.INFO]}"
    )
    _prepare_sheet(
        sheet,
        "Аудит метаданных",
        subtitle,
        headers,
        tab_color=CORAL_DARK,
    )
    severity_fills = {
        AuditSeverity.ERROR: CORAL,
        AuditSeverity.WARNING: BLUE,
        AuditSeverity.INFO: GREEN,
    }
    severity_fonts = {
        AuditSeverity.ERROR: Font(color=CORAL_DARK, bold=True, size=9),
        AuditSeverity.WARNING: Font(color=BLUE_DARK, bold=True, size=9),
        AuditSeverity.INFO: Font(color=GREEN_DARK, bold=True, size=9),
    }
    fix_font = Font(color=GREEN_DARK, bold=True, size=9)
    for offset, issue in enumerate(issues):
        row_number = DATA_START_ROW + offset
        track = tracks_by_id.get(issue.track_id)
        values = [
            offset + 1,
            issue.severity.value,
            issue.category.value,
            issue.field_label,
            issue.message,
            issue.current_value,
            issue.suggested_value,
            issue.confidence.value,
            "Да, после подтверждения" if issue.safe_fix else "Нет",
            track.relative_folder if track else "",
            track.filename if track else "",
            str(track.path) if track else "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, _safe(value))
        _style_data_row(
            sheet,
            row_number,
            len(headers),
            severity_fills[issue.severity],
        )
        sheet.cell(row_number, 2).font = severity_fonts[issue.severity]
        if issue.safe_fix:
            fix_cell = sheet.cell(row_number, 9)
            fix_cell.fill = _fill(GREEN)
            fix_cell.font = fix_font
    _fit_columns(sheet, maximum=60, minimum=11)
    sheet.column_dimensions["A"].width = 7
    _finish_table(sheet)


def export_xlsx(
    tracks: list[Track],
    destination: Path,
    selected_sheets: dict[str, bool],
    selected_columns: list[str],
    colors: dict[str, str],
    audit_issues: list[MetadataAuditIssue] | None = None,
) -> Path:
    if not tracks:
        raise ValueError("Нет данных для экспорта")
    fields = [field for field in selected_columns if field in EXPORT_COLUMNS]
    if not fields:
        fields = ["base_title", "status"]
    if not any(selected_sheets.values()):
        selected_sheets = dict(selected_sheets)
        selected_sheets["all"] = True

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Отчёт Discography Analyzer"
    workbook.properties.subject = "Анализ локальной музыкальной коллекции"
    workbook.properties.creator = "Discography Analyzer"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    if selected_sheets.get("all"):
        _write_track_sheet(workbook, "Все треки", tracks, fields, colors)
    if selected_sheets.get("attention"):
        attention = [
            track
            for track in tracks
            if track.status
            in (TrackStatus.DUPLICATE, TrackStatus.VARIATION, TrackStatus.CONFLICT)
        ]
        if attention:
            _write_track_sheet(
                workbook, "Только дубликаты", attention, fields, colors
            )
    if selected_sheets.get("matrix"):
        _write_matrix(workbook, tracks, colors)
    if selected_sheets.get("folders"):
        _write_folder_summary(workbook, tracks, colors)
    if selected_sheets.get("audit"):
        _write_audit_sheet(
            workbook,
            list(audit_issues or []),
            {track.id: track for track in tracks},
        )

    if not workbook.sheetnames:
        _write_track_sheet(workbook, "Все треки", tracks, fields, colors)
    destination = destination.with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(
            dir=destination.parent,
            prefix=f".{destination.stem}-",
            suffix=".xlsx",
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        temporary.replace(destination)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return destination
