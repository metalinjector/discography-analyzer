import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import PropertyMock, patch

from openpyxl import load_workbook

from discography_analyzer.audit_models import (
    AuditCategory,
    AuditConfidence,
    AuditFixKind,
    AuditSeverity,
    MetadataAuditIssue,
)
from discography_analyzer.exporter import export_xlsx
from discography_analyzer.models import Track, TrackStatus


def track(identifier: str, status: TrackStatus, folder: str) -> Track:
    path = Path(f"/music/{folder}/{identifier}.mp3")
    return Track(
        id=identifier,
        path=path,
        root=Path("/music"),
        filename=path.name,
        relative_path=f"{folder}/{path.name}",
        relative_folder=folder,
        format="MP3",
        size_bytes=1024 * 1024,
        filename_title="Song",
        comparison_title="Song",
        base_title="Song",
        status=status,
        title="Song",
        artist="Artist A",
        album="Album",
        duration_seconds=180,
    )


class ExporterTests(unittest.TestCase):
    def test_export_creates_requested_sheets_and_total(self) -> None:
        tracks = [
            track("one", TrackStatus.DUPLICATE, "Album 1"),
            track("two", TrackStatus.DUPLICATE, "Album 2"),
        ]
        with TemporaryDirectory() as directory:
            path = export_xlsx(
                tracks,
                Path(directory) / "report",
                {"all": True, "attention": True, "matrix": True, "folders": True},
                ["base_title", "status", "relative_folder"],
                {
                    "Уникальный": "E2F0D9",
                    "Дубликат": "FFF2CC",
                    "Вариация": "DDEBF7",
                    "Конфликт": "FCE4D6",
                },
            )
            workbook = load_workbook(path)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Все треки",
                    "Только дубликаты",
                    "Матрица дубликатов",
                    "Сводка по папкам",
                ],
            )
            summary = workbook["Сводка по папкам"]
            self.assertEqual(summary.cell(summary.max_row, 1).value, "ИТОГО")
            self.assertEqual(summary.freeze_panes, "A5")
            self.assertIn(
                "Проверка комплекта",
                [cell.value for cell in summary[4]],
            )
            self.assertEqual(summary.row_dimensions[5].height, 17.25)
            self.assertEqual(summary.sheet_view.showGridLines, False)
            all_tracks = workbook["Все треки"]
            self.assertEqual(all_tracks.cell(1, 1).value, "Все треки")
            self.assertEqual(all_tracks.row_dimensions[5].height, 17.25)
            self.assertEqual(all_tracks.freeze_panes, "A5")

    def test_disabling_all_sheets_still_creates_all_tracks(self) -> None:
        with TemporaryDirectory() as directory:
            path = export_xlsx(
                [track("one", TrackStatus.UNIQUE, "Album")],
                Path(directory) / "fallback.xlsx",
                {"all": False, "attention": False, "matrix": False, "folders": False},
                [],
                {},
            )
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["Все треки"])

    def test_empty_metadata_fields_do_not_break_column_sizing(self) -> None:
        empty_track = track("empty", TrackStatus.UNIQUE, "Album")
        empty_track.artist = ""
        empty_track.album = ""
        empty_track.genre = ""
        empty_track.comment = ""
        with TemporaryDirectory() as directory:
            path = export_xlsx(
                [empty_track],
                Path(directory) / "empty-metadata.xlsx",
                {"all": True},
                ["artist", "album", "genre", "comment"],
                {},
            )

            workbook = load_workbook(path)
            sheet = workbook["Все треки"]
            self.assertEqual(sheet.cell(5, 2).value, None)
            self.assertEqual(sheet.row_dimensions[5].height, 17.25)

    def test_leading_spaces_do_not_bypass_formula_protection(self) -> None:
        unsafe_track = track("unsafe", TrackStatus.UNIQUE, "Album")
        unsafe_track.artist = "  =HYPERLINK(\"https://example.invalid\")"
        with TemporaryDirectory() as directory:
            path = export_xlsx(
                [unsafe_track],
                Path(directory) / "safe.xlsx",
                {"all": True},
                ["artist"],
                {},
            )

            workbook = load_workbook(path, data_only=False)
            value = workbook["Все треки"].cell(5, 2).value
            self.assertEqual(
                value,
                "'  =HYPERLINK(\"https://example.invalid\")",
            )

    def test_control_characters_do_not_abort_export(self) -> None:
        dirty_track = track("dirty", TrackStatus.UNIQUE, "Album")
        dirty_track.title = "Song\x01Name"
        # Removing the character must not let a formula slip through.
        dirty_track.artist = "\x0c=1+1"
        with TemporaryDirectory() as directory:
            path = export_xlsx(
                [dirty_track],
                Path(directory) / "dirty.xlsx",
                {"all": True},
                ["title", "artist"],
                {},
            )

            sheet = load_workbook(path)["Все треки"]
            self.assertEqual(sheet.cell(5, 2).value, "SongName")
            self.assertEqual(sheet.cell(5, 3).value, "'=1+1")

    def test_failed_export_preserves_existing_report(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            path.write_bytes(b"existing report")

            with (
                patch(
                    "discography_analyzer.exporter.Workbook.save",
                    side_effect=OSError("disk write failed"),
                ),
                self.assertRaisesRegex(OSError, "disk write failed"),
            ):
                export_xlsx(
                    [track("one", TrackStatus.UNIQUE, "Album")],
                    path,
                    {"all": True},
                    ["artist"],
                    {},
                )

            self.assertEqual(path.read_bytes(), b"existing report")
            self.assertEqual(
                list(Path(directory).glob(".report-*.xlsx")),
                [],
            )

    def test_row_styling_does_not_scan_whole_sheet_per_row(self) -> None:
        # ``sheet[row]`` asks openpyxl for ``max_column`` which walks every
        # cell, turning a 50k-track report into an O(rows²) export.
        tracks = [
            track(f"t{index}", TrackStatus.UNIQUE, "Album") for index in range(40)
        ]
        with (
            TemporaryDirectory() as directory,
            patch(
                "openpyxl.worksheet.worksheet.Worksheet.max_column",
                new_callable=PropertyMock,
                return_value=5,
            ) as max_column,
        ):
            export_xlsx(
                tracks,
                Path(directory) / "scan.xlsx",
                {"all": True},
                ["base_title", "status", "relative_folder", "artist"],
                {},
            )

        # A handful of calls for auto-filter/column fitting is fine; one call
        # per data row is the regression we guard against.
        self.assertLess(max_column.call_count, len(tracks))

    def test_audit_sheet_contains_issue_and_safe_fix_marker(self) -> None:
        item = track("one", TrackStatus.UNIQUE, "Album")
        issue = MetadataAuditIssue(
            id="audit-one",
            track_id=item.id,
            severity=AuditSeverity.WARNING,
            category=AuditCategory.ARTISTS,
            field="artist",
            message="Неверный разделитель",
            current_value="Riana;Moby",
            suggested_value="Riana, Moby",
            fix_kind=AuditFixKind.NORMALIZE_ARTISTS,
            confidence=AuditConfidence.HIGH,
        )
        with TemporaryDirectory() as directory:
            path = export_xlsx(
                [item],
                Path(directory) / "audit.xlsx",
                {"all": False, "audit": True},
                ["artist"],
                {},
                [issue],
            )
            workbook = load_workbook(path)

        self.assertEqual(workbook.sheetnames, ["Аудит метаданных"])
        sheet = workbook["Аудит метаданных"]
        self.assertEqual(sheet.cell(5, 5).value, "Неверный разделитель")
        self.assertEqual(sheet.cell(5, 9).value, "Да, после подтверждения")
        self.assertEqual(sheet.row_dimensions[5].height, 17.25)


if __name__ == "__main__":
    unittest.main()
